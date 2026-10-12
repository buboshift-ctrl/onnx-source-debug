"""
admit/mlir-arx/tail 파이프라인 사이드바를, SVG 삽입이 아니라 pptx 네이티브 도형(사각형+텍스트)
으로 그려서 슬라이드 1장짜리 .pptx로 만든다.

Keynote는 .key(애플 비공개 바이너리 포맷)를 직접 쓸 수 있는 라이브러리가 없어서, 대신
"Keynote가 File > Open 으로 그대로 열어서 자동 변환하는" .pptx를 만든다. SVG를 이미지로
삽입하는 방식 대신 네이티브 도형으로 그리는 이유: 변환 후에도 도형/텍스트가 계속 편집
가능한 상태로 들어가고, 한글도 실제 텍스트 런(run)이라 폰트 렌더링 문제가 없다.

실행: python3 make_pipeline_keynote_slide.py
출력: pipeline_sidebar_keynote.pptx (같은 폴더)
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

HERE = os.path.dirname(os.path.abspath(__file__))

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
ACCENT = RGBColor(0x2E, 0x6F, 0x95)      # admit/tail 공통 "우리 소유" 색
LOCKED = RGBColor(0x8A, 0x6D, 0x3B)      # mlir-arx "못 건드림"
GRAY = RGBColor(0x5A, 0x5F, 0x66)
ENDPOINT_FILL = RGBColor(0xEE, 0xF1, 0xF4)
ENDPOINT_BORDER = RGBColor(0xC7, 0xCC, 0xD1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SUBTEXT_ON_COLOR = RGBColor(0xE9, 0xF0, 0xF5)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

STAGES = [
    {"id": "onnx", "kind": "endpoint", "h": Inches(0.6),
     "title": "ONNX 모델", "sub": "입력"},
    {"id": "admit", "kind": "ours", "h": Inches(1.0),
     "title": "admit 단계", "sub": "ONNX 로드 직후", "tag": "우리가 만짐"},
    {"id": "mlirarx", "kind": "locked", "h": Inches(1.0),
     "title": "mlir-arx (PNU)", "sub": "onnx → harx → EmitC", "tag": "절대 안 건드림"},
    {"id": "tail", "kind": "ours", "h": Inches(1.0),
     "title": "tail 단계", "sub": "EmitC → C 생성 이후", "tag": "우리 소유"},
    {"id": "artifact", "kind": "endpoint", "h": Inches(0.6),
     "title": "host-mode 실행 아티팩트", "sub": "출력"},
]

BOX_W = Inches(3.1)
GAP = Inches(0.32)
CONTENT_TOP = Inches(1.2)   # 제목 바로 아래 고정 시작 위치 (세로 중앙정렬로 인한 하단 캡션과의 겹침 방지)


def add_rect(slide, x, y, w, h, fill, border, border_w):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.adjustments[0] = 0.12
    shp.shadow.inherit = False
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = border
    shp.line.width = border_w
    return shp


def set_lines(shape, lines, anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    for i, (text, size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = "Pretendard"


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.shadow.inherit = False
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()

    # 제목
    title = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(9), Inches(0.6))
    tf = title.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "우리가 수정 가능한 지점은 두 곳뿐 — admit / tail"
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r.font.name = "Pretendard"

    start_x = Emu(int((SLIDE_W - BOX_W) / 2))
    start_y = CONTENT_TOP

    y = start_y
    box_geoms = []
    for stage in STAGES:
        box_geoms.append((y, stage["h"]))
        y = Emu(int(y + stage["h"] + GAP))

    # 화살표
    for i in range(len(STAGES) - 1):
        top_of_next = box_geoms[i + 1][0]
        bottom_of_this = Emu(int(box_geoms[i][0] + box_geoms[i][1]))
        cx = Emu(int(start_x + BOX_W / 2))
        conn = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, cx, bottom_of_this, cx, top_of_next
        )
        conn.line.color.rgb = GRAY
        conn.line.width = Pt(1.75)

    # 박스
    for stage, (by, bh) in zip(STAGES, box_geoms):
        kind = stage["id"], stage["kind"]
        if stage["kind"] == "endpoint":
            box = add_rect(slide, start_x, by, BOX_W, bh,
                            fill=ENDPOINT_FILL, border=ENDPOINT_BORDER, border_w=Pt(1.25))
            set_lines(box, [
                (stage["title"], 13, True, GRAY),
                (stage["sub"], 10.5, False, GRAY),
            ])
        else:
            main = ACCENT if stage["kind"] == "ours" else LOCKED
            box = add_rect(slide, start_x, by, BOX_W, bh,
                            fill=main, border=NAVY, border_w=Pt(1.75))
            set_lines(box, [
                (stage["title"], 16, True, WHITE),
                (stage["sub"], 11.5, False, SUBTEXT_ON_COLOR),
                (f'[{stage["tag"]}]', 10, True, SUBTEXT_ON_COLOR),
            ])

    footer = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(11), Inches(0.35))
    ftf = footer.text_frame
    fp = ftf.paragraphs[0]
    fr = fp.add_run()
    fr.text = "이 도형/텍스트는 Keynote에서 열어도 그대로 편집 가능합니다 (File > Open으로 자동 변환)."
    fr.font.size = Pt(10)
    fr.font.color.rgb = GRAY
    fr.font.name = "Pretendard"

    out = os.path.join(HERE, "pipeline_sidebar_keynote.pptx")
    prs.save(out)
    print("저장:", out)


if __name__ == "__main__":
    build()
